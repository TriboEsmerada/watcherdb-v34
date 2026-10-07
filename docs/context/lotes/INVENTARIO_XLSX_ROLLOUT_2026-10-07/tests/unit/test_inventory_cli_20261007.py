# -*- coding: utf-8 -*-
"""2026-10-07 -- INVENTARIO_XLSX_ROLLOUT (lote B): .xlsx, servers.json canonico, politica Windows Auth, rollout no
formato do pool, scripts de grants por instancia (SID so' por AG, sem password, sem EXECUTE AS), modelo xlsx."""
from __future__ import annotations

import json
import re
import threading
from pathlib import Path
from types import SimpleNamespace

import pytest

from watcherdb.install import inventory as inv
from watcherdb.install import cli

ROOT = Path(__file__).resolve().parents[2]
GRANTS_DOC = ROOT / "docs" / "security" / "GRANTS_SQL_MONITORING_INSTANCIA.sql"

CAB = list(inv.COLUNAS)


def _xlsx(tmp_path, linhas, folha="Inventario", cabecalho=None):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = folha
    ws.append(cabecalho or CAB)
    for l in linhas:
        ws.append(l)
    p = tmp_path / "inv.xlsx"
    wb.save(str(p))
    return p


# ------------------------------------------------------------------ xlsx
def test_xlsx_le_numeros_booleanos_e_cabecalhos_com_espacos(tmp_path):
    p = _xlsx(tmp_path, [
        ["SQL01.banco.pt", "MSSQLSERVER", 1433.0, "sql", "", "Core", "production", 1.0, True, False, None, None],
        ["sql02", "PROD", 1434, "SQL", None, "Risco", "Production", 2, "sim", "true", "AG_Fin", "LST"],
        [None] * 12,
    ], cabecalho=["Host", "Instance ", "port", "Auth Mode", "username", "description", "environment", "priority", "enabled", "has alwayson", "ag_name", "ag_listener"])
    entries, issues = inv.parse_xlsx(p, "watcherdb")
    assert [i for i in issues if i.severity == "error"] == []
    assert len(entries) == 2
    assert entries[0].port == 1433 and entries[0].priority == 1 and entries[0].enabled is True
    assert entries[1].id == "SQL02_PROD" and entries[1].has_alwayson is True and entries[1].ag_name == "AG_Fin"
    assert entries[1].username == "watcherdb" and entries[1].password == inv.PASSWORD_PLACEHOLDER


def test_xlsx_folha_inventario_tem_prioridade_e_formula_sem_cache_da_erro(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Outra"
    ws0.append(["lixo"])
    ws = wb.create_sheet("Inventario")
    ws.append(CAB)
    ws.append(["=A9", "MSSQLSERVER", 1433])
    ws.append(["SQL03", "I01", 1433])
    p = tmp_path / "f.xlsx"
    wb.save(str(p))
    entries, issues = inv.parse_xlsx(p, "watcherdb")
    assert len(entries) == 1 and entries[0].id == "SQL03_I01"
    assert any("formula" in i.message for i in issues if i.severity == "error")


def test_xlsx_sem_coluna_obrigatoria(tmp_path):
    p = _xlsx(tmp_path, [["SQL01", "MSSQLSERVER"]], cabecalho=["host", "instance"])
    entries, issues = inv.parse_xlsx(p, "x")
    assert entries == [] and issues[0].field == "header" and "port" in issues[0].message


def test_ler_inventario_detecta_formatos(tmp_path):
    p = _xlsx(tmp_path, [["SQL01", "MSSQLSERVER", 1433]])
    assert inv.ler_inventario(p, "x")[3] == "xlsx"
    j = tmp_path / "a.json"
    j.write_text(json.dumps([{"host": "SQL01", "instance": "I01", "port": 1433}]), encoding="utf-8")
    assert inv.ler_inventario(j, "x")[3] == "json"
    c = tmp_path / "a.json2.json"
    c.write_text(json.dumps({"master_server": {"host": "M"}, "monitored_servers": [{"id": "SQL01_I01", "host": "SQL01", "instance": "I01", "port": 1433, "use_windows_auth": False, "username": "u", "password": "encrypted:abc", "databases": ["x"]}]}), encoding="utf-8")
    e, i, m, f = inv.ler_inventario(c, "x")
    assert f == "servers.json" and m == {"host": "M"} and e[0].id == "SQL01_I01" and e[0].password == "encrypted:abc"
    assert inv.ler_inventario(tmp_path / "a.txt", "x")[3] == ".txt"


# ------------------------------------------------------------------ politica e rollout
def _entries():
    return [
        inv.ServerEntry("SQL01_I01", "sql01", "I01", 1433, "", "production", 1, True, False, "watcherdb", inv.PASSWORD_PLACEHOLDER, "SQL Server", False, None, None),
        inv.ServerEntry("SQL02_MSSQLSERVER", "sql02", "MSSQLSERVER", 1433, "", "production", 1, True, True, "", "", "SQL Server", False, None, None),
        inv.ServerEntry("SQL03_AG", "sql03", "AG", 1433, "", "production", 1, False, False, "watcherdb", inv.PASSWORD_PLACEHOLDER, "SQL Server", True, "AG_Fin", "LST"),
        inv.ServerEntry("SQL04_AG", "sql04", "AG", 1433, "", "production", 1, True, False, "watcherdb", inv.PASSWORD_PLACEHOLDER, "SQL Server", True, "AG_Fin", "LST"),
        inv.ServerEntry("SQL05_AG2", "sql05", "AG2", 1433, "", "production", 1, True, False, "watcherdb", inv.PASSWORD_PLACEHOLDER, "SQL Server", True, "AG_Outro", None),
    ]


def test_windows_auth_e_erro_salvo_opt_in():
    e = _entries()
    erros = inv.politica_windows_auth(e, allow_windows_auth=False)
    assert len(erros) == 1 and "SQL02_MSSQLSERVER" in erros[0].message and "Regra de Ouro #2" in erros[0].message
    assert inv.politica_windows_auth(e, allow_windows_auth=True) == []


def test_rollout_so_sql_auth_enabled_maiusculas_e_verificados():
    e = _entries()
    r = inv.gerar_rollout(e)
    assert r["sql_auth_servers"] == ["SQL01_I01", "SQL04_AG", "SQL05_AG2"]   # sem windows, sem disabled
    r2 = inv.gerar_rollout(e, verificados=["sql04_ag"])
    assert r2["sql_auth_servers"] == ["SQL04_AG"] and r2["_gerado"]["origem"] == "preflight-fleet"


def test_rollout_round_trip_com_o_loader_do_pool(tmp_path):
    from api.connection_pool import SQLServerConnectionPool
    p = tmp_path / "sql_auth_rollout.json"
    inv.write_rollout_json(p, inv.gerar_rollout(_entries()))
    falso = SimpleNamespace(_sql_auth_rollout_lock=threading.Lock(), _sql_auth_rollout_path=str(p),
                            _sql_auth_rollout=set(), _sql_auth_rollout_loaded_at=0.0)
    lidos = SQLServerConnectionPool._load_sql_auth_rollout(falso)
    assert lidos == {"SQL01_I01", "SQL04_AG", "SQL05_AG2"}


# ------------------------------------------------------------------ grants por instancia
def test_corpo_dos_grants_igual_ao_ficheiro_canonico():
    """Anti-drift: o corpo embutido tem de ser o do docs/security (sem cabecalho, sem guarda, sem EXECUTE AS)."""
    txt = GRANTS_DOC.read_text(encoding="utf-8").replace("\r\n", "\n")
    ini = txt.index("-- ---- servidor ")
    fim = txt.index("-- ---- verificacao")
    assert txt[ini:fim].strip() == inv.GRANTS_CORPO.strip()
    assert "EXECUTE AS" not in inv.GRANTS_CORPO


def test_scripts_por_instancia_login_substituido_sid_so_por_ag_sem_password(tmp_path):
    e = _entries()
    out = tmp_path / "grants"
    res = inv.escrever_scripts_grants(e, "watcherdb", out)
    ids = [i["id"] for i in res["instancias"]]
    assert ids == ["SQL01_I01", "SQL04_AG", "SQL05_AG2"]        # sem windows (SQL02), sem disabled (SQL03)
    assert set(res["sids"]) == {"AG_Fin", "AG_Outro"} and res["sids"]["AG_Fin"] != res["sids"]["AG_Outro"]
    assert re.fullmatch(r"0x[0-9A-F]{32}", res["sids"]["AG_Fin"])
    g1 = (out / "grants_SQL01_I01.sql").read_text(encoding="utf-8")
    g4 = (out / "grants_SQL04_AG.sql").read_text(encoding="utf-8")
    assert "sql_monitoring" not in g1 and "[watcherdb]" in g1 and "TO [watcherdb]" in g1
    assert "SID =" not in g1 and f"SID = {res['sids']['AG_Fin']}" in g4
    assert inv.PLACEHOLDER_SCRIPT in g1 and "RAISERROR('Password por preencher" in g1
    assert "EXECUTE AS" not in g1 and "CHECK_POLICY = ON" in g1 and "CHECK_EXPIRATION = OFF" in g1
    assert "SO DE LEITURA" in g1 and "ROLLBACK: ficheiro rollback_SQL01_I01.sql" in g1
    r1 = (out / "rollback_SQL01_I01.sql").read_text(encoding="utf-8")
    assert "DROP LOGIN [watcherdb]" in r1 and "DROP USER [watcherdb]" in r1 and "nao matar sessoes" in r1
    idx = (out / "INDICE.md").read_text(encoding="utf-8")
    assert "SQL04_AG" in idx and res["instancias"][0]["grants_sha256"][:16] in idx
    assert json.loads((out / "sids.json").read_text(encoding="utf-8")) == res["sids"]


def test_sids_reutilizados_entre_execucoes(tmp_path):
    e = _entries()
    r1 = inv.escrever_scripts_grants(e, "watcherdb", tmp_path / "g")
    r2 = inv.escrever_scripts_grants(e, "watcherdb", tmp_path / "g", sids=r1["sids"])
    assert r2["sids"] == r1["sids"]


def test_texto_livre_nao_quebra_comentarios_sql():
    row = {"host": "H", "instance": "I", "port": 1433, "description": "linha1\nlinha2 -- fim */ x", "ag_name": "AG\n--"}
    e, _ = inv._parse_row(row, 2, "watcherdb")
    assert "\n" not in e.description and "--" not in e.description and "*/" not in e.description
    assert "\n" not in e.ag_name and "--" not in e.ag_name


# ------------------------------------------------------------------ modelo
def test_modelo_xlsx_gera_e_reparse(tmp_path):
    p = inv.gerar_modelo_xlsx(tmp_path / "modelo.xlsx")
    entries, issues = inv.parse_xlsx(p, "watcherdb")
    assert len(entries) == 1 and entries[0].host == "SQL01.exemplo.local" and entries[0].port == 1433
    from openpyxl import load_workbook
    wb = load_workbook(str(p))
    assert wb.sheetnames == ["Inventario", "Instrucoes"]
    assert [c.value for c in wb["Inventario"][1]] == CAB


# ------------------------------------------------------------------ cli
def test_cli_inventory_fim_a_fim(tmp_path, capsys):
    p = _xlsx(tmp_path, [
        ["SQL01", "I01", 1433, "sql", "", "a", "production", 1, True, False, "", ""],
        ["SQL04", "AG", 1433, "sql", "", "b", "production", 1, True, True, "AG_Fin", "LST"],
    ])
    out = tmp_path / "cfg"
    rc = cli.main(["inventory", "--input", str(p), "--output-dir", str(out), "--master-host", "M", "--login", "watcherdb", "--rollout-sem-verificacao"])
    assert rc == 0, capsys.readouterr().err
    doc = json.loads((out / "servers.json").read_text(encoding="utf-8"))
    assert doc["master_server"]["database"] == "WatcherDB" and doc["master_server"]["username"] == "watcherdb"
    assert [s["id"] for s in doc["monitored_servers"]] == ["SQL01_I01", "SQL04_AG"]
    assert (out / "grants" / "grants_SQL04_AG.sql").exists() and (out / "grants" / "INDICE.md").exists()
    assert json.loads((out / "sql_auth_rollout.json").read_text(encoding="utf-8"))["sql_auth_servers"] == ["SQL01_I01", "SQL04_AG"]


def test_cli_inventory_windows_auth_recusado_sem_opt_in(tmp_path, capsys):
    p = _xlsx(tmp_path, [["SQL02", "MSSQLSERVER", 1433, "windows"]])
    assert cli.main(["inventory", "--input", str(p), "--output-dir", str(tmp_path / "o"), "--master-host", "M"]) == 1
    assert "Regra de Ouro #2" in capsys.readouterr().err
    assert cli.main(["inventory", "--input", str(p), "--output-dir", str(tmp_path / "o2"), "--master-host", "M", "--allow-windows-auth"]) == 0


def test_cli_inventory_sem_rollout_por_omissao(tmp_path, capsys):
    p = _xlsx(tmp_path, [["SQL01", "I01", 1433]])
    out = tmp_path / "o"
    assert cli.main(["inventory", "--input", str(p), "--output-dir", str(out), "--master-host", "M"]) == 0
    assert not (out / "sql_auth_rollout.json").exists()
    assert "preflight-fleet" in capsys.readouterr().out


def test_cli_inventory_template(tmp_path):
    assert cli.main(["inventory-template", "--output", str(tmp_path / "m.xlsx")]) == 0
    assert (tmp_path / "m.xlsx").exists()


def test_cli_install_help_lista_os_novos(capsys):
    assert cli.main(["install-help"]) == 0
    out = capsys.readouterr().out
    assert "inventory" in out and "inventory-template" in out


def test_shim_em_deploy_reexporta_a_api_antiga():
    import importlib, sys as _sys
    d = str(ROOT / "deploy")
    if d not in _sys.path:
        _sys.path.insert(0, d)
    fip = importlib.import_module("farm_inventory_parser")
    assert fip.parse_csv is inv.parse_csv and fip.main is inv.main and fip.PASSWORD_PLACEHOLDER == inv.PASSWORD_PLACEHOLDER


def test_dispatcher_do_servico_inclui_inventory():
    svc = (ROOT / "watcherdb_service.py").read_text(encoding="utf-8")
    assert '"inventory", "inventory-template"' in svc
